from datetime import datetime, date, timedelta
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from app.database import get_db
from app.auth import get_current_user, require_admin_or_technician
from app.models import (
    User, Greenhouse, Zone, Device, Sensor, SensorReading,
    Alarm, AlarmStatus, Task, TaskStatus, PlantBatch,
    ProductionRecord, QualityGrade
)
from app.schemas import DashboardStats

router = APIRouter(tags=["报表与统计"])


@router.get("/api/dashboard/stats", response_model=DashboardStats)
async def get_dashboard_stats(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    total_greenhouses = db.query(Greenhouse).filter(
        Greenhouse.is_active == True
    ).count()
    total_zones = db.query(Zone).filter(
        Zone.is_active == True
    ).count()
    total_devices = db.query(Device).filter(
        Device.is_active == True
    ).count()
    active_alarms = db.query(Alarm).filter(
        Alarm.status.in_([AlarmStatus.PENDING, AlarmStatus.ACKNOWLEDGED])
    ).count()
    pending_tasks = db.query(Task).filter(
        Task.status == TaskStatus.PENDING
    ).count()
    return DashboardStats(
        total_greenhouses=total_greenhouses,
        total_zones=total_zones,
        total_devices=total_devices,
        active_alarms=active_alarms,
        pending_tasks=pending_tasks
    )


@router.get("/api/dashboard/environment")
async def get_environment_overview(
    greenhouse_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = db.query(Sensor).filter(Sensor.is_active == True)
    if greenhouse_id:
        query = query.join(Zone).filter(Zone.greenhouse_id == greenhouse_id)
    sensors = query.all()
    overview = []
    for sensor in sensors:
        latest_reading = db.query(SensorReading).filter(
            SensorReading.sensor_id == sensor.id
        ).order_by(SensorReading.reading_time.desc()).first()
        overview.append({
            "sensor_id": sensor.id,
            "sensor_name": sensor.name,
            "sensor_type": sensor.sensor_type.value,
            "zone_name": sensor.zone.name if sensor.zone else None,
            "current_value": latest_reading.value if latest_reading else None,
            "unit": sensor.unit,
            "reading_time": latest_reading.reading_time.isoformat() if latest_reading else None
        })
    return overview


@router.get("/api/dashboard/devices")
async def get_devices_overview(
    greenhouse_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    query = db.query(Device).filter(Device.is_active == True)
    if greenhouse_id:
        query = query.join(Zone).filter(Zone.greenhouse_id == greenhouse_id)
    devices = query.all()
    overview = []
    for device in devices:
        overview.append({
            "device_id": device.id,
            "device_name": device.name,
            "device_type": device.device_type.value,
            "zone_name": device.zone.name if device.zone else None,
            "is_on": device.is_on,
            "is_auto_mode": device.is_auto_mode
        })
    return overview


@router.get("/api/dashboard/recent-alarms")
async def get_recent_alarms(
    limit: int = Query(10, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    alarms = db.query(Alarm).order_by(
        Alarm.created_at.desc()
    ).limit(limit).all()
    return alarms


@router.get("/api/dashboard/today-tasks")
async def get_today_tasks(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    today = date.today()
    query = db.query(Task).filter(Task.scheduled_date == today)
    if current_user.role != "admin":
        query = query.filter(
            (Task.assigned_user_id == current_user.id) |
            (Task.assigned_user_id.is_(None))
        )
    tasks = query.order_by(Task.scheduled_time.asc().nullslast()).all()
    return tasks


@router.get("/api/reports/production")
async def get_production_report(
    greenhouse_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    query = db.query(ProductionRecord).filter(
        ProductionRecord.harvest_date >= start_date,
        ProductionRecord.harvest_date <= end_date
    )
    if greenhouse_id:
        query = query.filter(ProductionRecord.greenhouse_id == greenhouse_id)
    records = query.all()
    total_quantity = sum(r.quantity for r in records)
    total_yield = sum(r.yield_amount or 0 for r in records)
    quality_distribution = {
        "excellent": 0,
        "good": 0,
        "fair": 0,
        "poor": 0
    }
    for r in records:
        grade = r.quality_grade.value if hasattr(r.quality_grade, "value") else r.quality_grade
        if grade in quality_distribution:
            quality_distribution[grade] += r.quantity
    by_variety = {}
    for r in records:
        if r.flower_variety not in by_variety:
            by_variety[r.flower_variety] = {
                "quantity": 0,
                "yield_amount": 0
            }
        by_variety[r.flower_variety]["quantity"] += r.quantity
        by_variety[r.flower_variety]["yield_amount"] += r.yield_amount or 0
    return {
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "total_records": len(records),
        "total_quantity": total_quantity,
        "total_yield": total_yield,
        "avg_yield_per_unit": total_yield / total_quantity if total_quantity > 0 else 0,
        "quality_distribution": quality_distribution,
        "by_variety": by_variety
    }


@router.get("/api/reports/production/export")
async def export_production_report(
    greenhouse_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin_or_technician())
):
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    query = db.query(ProductionRecord).filter(
        ProductionRecord.harvest_date >= start_date,
        ProductionRecord.harvest_date <= end_date
    )
    if greenhouse_id:
        query = query.filter(ProductionRecord.greenhouse_id == greenhouse_id)
    records = query.order_by(ProductionRecord.harvest_date.asc()).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "生产报表"
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    headers = ["ID", "温室", "区域", "花卉品种", "数量", "品质等级", "收获日期", "产量", "备注"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    quality_map = {
        "excellent": "优秀",
        "good": "良好",
        "fair": "一般",
        "poor": "较差"
    }
    for row_idx, record in enumerate(records, 2):
        greenhouse = db.query(Greenhouse).filter(
            Greenhouse.id == record.greenhouse_id
        ).first()
        zone = None
        if record.zone_id:
            zone = db.query(Zone).filter(Zone.id == record.zone_id).first()
        grade = record.quality_grade.value if hasattr(record.quality_grade, "value") else record.quality_grade
        data = [
            record.id,
            greenhouse.name if greenhouse else "",
            zone.name if zone else "",
            record.flower_variety,
            record.quantity,
            quality_map.get(grade, grade),
            record.harvest_date.isoformat(),
            record.yield_amount or 0,
            record.notes or ""
        ]
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
    for col in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col) if col <= 26 else chr(64 + col // 26) + chr(64 + col % 26)].width = 15
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"production_report_{start_date}_{end_date}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename={filename}"
        }
    )


@router.get("/api/reports/environment")
async def get_environment_report(
    greenhouse_id: Optional[int] = None,
    zone_id: Optional[int] = None,
    sensor_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=7)
    query = db.query(Sensor).filter(Sensor.is_active == True)
    if sensor_id:
        query = query.filter(Sensor.id == sensor_id)
    elif zone_id:
        query = query.filter(Sensor.zone_id == zone_id)
    elif greenhouse_id:
        query = query.join(Zone).filter(Zone.greenhouse_id == greenhouse_id)
    sensors = query.all()
    report = []
    for sensor in sensors:
        readings = db.query(SensorReading).filter(
            SensorReading.sensor_id == sensor.id,
            SensorReading.reading_time >= start_date,
            SensorReading.reading_time <= end_date
        ).order_by(SensorReading.reading_time.asc()).all()
        if readings:
            values = [r.value for r in readings]
            report.append({
                "sensor_id": sensor.id,
                "sensor_name": sensor.name,
                "sensor_type": sensor.sensor_type.value,
                "unit": sensor.unit,
                "zone_name": sensor.zone.name if sensor.zone else None,
                "avg_value": sum(values) / len(values),
                "min_value": min(values),
                "max_value": max(values),
                "reading_count": len(readings),
                "readings": [
                    {"time": r.reading_time.isoformat(), "value": r.value}
                    for r in readings
                ]
            })
    return report


@router.get("/api/reports/tasks")
async def get_task_report(
    assigned_user_id: Optional[int] = None,
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    if not end_date:
        end_date = date.today()
    if not start_date:
        start_date = end_date - timedelta(days=30)
    query = db.query(Task).filter(
        Task.scheduled_date >= start_date,
        Task.scheduled_date <= end_date
    )
    if assigned_user_id:
        query = query.filter(Task.assigned_user_id == assigned_user_id)
    if current_user.role != "admin":
        query = query.filter(Task.assigned_user_id == current_user.id)
    tasks = query.all()
    status_distribution = {
        "pending": 0,
        "in_progress": 0,
        "completed": 0,
        "cancelled": 0
    }
    type_distribution = {}
    for task in tasks:
        status = task.status.value if hasattr(task.status, "value") else task.status
        if status in status_distribution:
            status_distribution[status] += 1
        task_type = task.task_type.value if hasattr(task.task_type, "value") else task.task_type
        if task_type not in type_distribution:
            type_distribution[task_type] = 0
        type_distribution[task_type] += 1
    completed = status_distribution["completed"]
    total = len(tasks)
    return {
        "start_date": start_date.isoformat(),
        "end_date": end_date.isoformat(),
        "total_tasks": total,
        "status_distribution": status_distribution,
        "type_distribution": type_distribution,
        "completion_rate": completed / total if total > 0 else 0
    }
